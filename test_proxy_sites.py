#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
test_proxy_sites.py —— GitHub 文件下载代理站点测速与可用性测试（用户手动运行，不接入自动化）

功能：
  1. 抓取以下三个网站公开展示的 GitHub 下载代理/镜像站点列表：
       - https://github.qqday.com/
       - https://github.akams.cn/
       - https://www.moretools.app/zh-CN/github-proxy
  2. 按域名去重；
  3. 通过每个代理下载同一个 GitHub 大文件（默认截取前 20MB），
     测量首字节延迟与下载速度；超时/失败的站点跳过并记录状态；
  4. 结果输出为 Excel（.xlsx）表格。

依赖：requests、openpyxl（pip install requests openpyxl）

用法：
  python test_proxy_sites.py                      # 默认 20MB、8 线程
  python test_proxy_sites.py --size 10 --workers 4
  python test_proxy_sites.py --file "https://github.com/owner/repo/releases/download/x/big.zip"
  python test_proxy_sites.py -o result.xlsx
"""

import argparse
import re
import sys
import threading
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime
from urllib.parse import urlparse

import requests

# ---------------------------------------------------------------------------
# 配置
# ---------------------------------------------------------------------------

# 默认测速文件（GitHub Release 大文件，只截取前 --size MB）
DEFAULT_TEST_FILE = (
    'https://github.com/microsoft/terminal/releases/download/v1.22.10731.0/'
    'Microsoft.WindowsTerminal_1.22.10731.0_x64.zip'
)

CONNECT_TIMEOUT = 10      # 连接超时（秒）
READ_TIMEOUT = 15         # 单次读取超时（秒）
TOTAL_TIMEOUT = 120       # 单站整体耗时上限（秒），超过视为超时跳过

UA = ('Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 '
      '(KHTML, like Gecko) Chrome/126.0 Safari/537.36')

# 三个来源站点
SOURCES = {
    'qqday':     'https://github.qqday.com/',
    'akams':     'https://github.akams.cn/',
    'moretools': 'https://www.moretools.app/zh-CN/github-proxy',
}

# akams 节点列表（开源项目 hubporg/ghproxy-next 的 components/nodes.ts）
AKAMS_NODES_URLS = [
    'https://cdn.jsdelivr.net/gh/hubporg/ghproxy-next@main/components/nodes.ts',
    'https://raw.githubusercontent.com/hubporg/ghproxy-next/main/components/nodes.ts',
]

# 抓取 moretools 页面时需排除的非节点域名
MORETOOLS_EXCLUDE = {
    'www.moretools.app', 'moretools.app', 'github.com', 'raw.githubusercontent.com',
    'schema.org', 'www.w3.org', 'fonts.googleapis.com', 'fonts.gstatic.com',
}

_print_lock = threading.Lock()
_ssl_fallback_warned = threading.Event()


def log(msg):
    with _print_lock:
        print(msg, flush=True)


def http_get(url, **kw):
    """
    requests.get 封装：优先正常校验证书；遇到本机 TLS 拦截代理等导致的
    证书校验失败时，回退到不校验证书（仅本工具脚本使用）并重试一次。
    """
    kw.setdefault('headers', {'User-Agent': UA})
    try:
        return requests.get(url, **kw)
    except requests.exceptions.SSLError:
        if not _ssl_fallback_warned.is_set():
            _ssl_fallback_warned.set()
            log('[提示] 检测到 TLS 证书校验失败（可能为本机 TLS 拦截代理），后续请求将跳过证书校验')
        import urllib3
        urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)
        kw['verify'] = False
        return requests.get(url, **kw)


# ---------------------------------------------------------------------------
# 1. 抓取三个网站的站点列表
# ---------------------------------------------------------------------------

def _norm_host(url_or_host):
    """把 URL 或域名规范化为小写 hostname（去协议、路径、端口、末尾点）。"""
    s = url_or_host.strip()
    if not s:
        return ''
    if '://' not in s:
        s = 'https://' + s
    try:
        host = urlparse(s).hostname or ''
    except ValueError:
        return ''
    return host.lower().rstrip('.')


def fetch_qqday():
    """github.qqday.com：节点在服务端渲染的 data-url 属性里。"""
    hosts = set()
    try:
        r = http_get(SOURCES['qqday'], timeout=CONNECT_TIMEOUT)
        r.encoding = 'utf-8'
        for m in re.finditer(r'data-url="(https?://[^"]+)"', r.text):
            h = _norm_host(m.group(1))
            if h:
                hosts.add(h)
        # 兜底：匹配节点名 span
        if not hosts:
            for m in re.finditer(r'node-pick__name">([^<]+)<', r.text):
                h = _norm_host(m.group(1))
                if h:
                    hosts.add(h)
    except Exception as e:
        log(f'[抓取] qqday 失败: {e}')
    return hosts


def fetch_akams():
    """github.akams.cn：节点列表在其开源仓库 components/nodes.ts 中。"""
    hosts = set()
    text = ''
    for url in AKAMS_NODES_URLS:
        try:
            r = http_get(url, timeout=CONNECT_TIMEOUT)
            if r.ok and 'value' in r.text:
                text = r.text
                break
        except Exception:
            continue
    if not text:
        log('[抓取] akams 失败: nodes.ts 所有镜像均不可达')
        return hosts
    for m in re.finditer(r'value:\s*"([^"]+)"', text):
        h = _norm_host(m.group(1))
        if h:
            hosts.add(h)
    return hosts


def fetch_moretools():
    """moretools.app：节点以完整 URL 文本形式列在页面中。"""
    hosts = set()
    try:
        r = http_get(SOURCES['moretools'], timeout=CONNECT_TIMEOUT)
        r.encoding = 'utf-8'
        for m in re.finditer(r'>(https?://[A-Za-z0-9.一-鿿-]+(?::\d+)?/?)<', r.text):
            h = _norm_host(m.group(1))
            if h and h not in MORETOOLS_EXCLUDE and not h.endswith('.moretools.app'):
                hosts.add(h)
    except Exception as e:
        log(f'[抓取] moretools 失败: {e}')
    return hosts


def collect_sites():
    """抓取三个来源，返回 {host: set(来源名)}，已按域名去重。"""
    site_sources = {}

    def merge(name, hosts):
        log(f'[抓取] {name}: {len(hosts)} 个站点')
        for h in hosts:
            site_sources.setdefault(h, set()).add(name)

    merge('qqday', fetch_qqday())
    merge('akams', fetch_akams())
    merge('moretools', fetch_moretools())
    log(f'[去重] 合计 {len(site_sources)} 个唯一站点')
    return site_sources


# ---------------------------------------------------------------------------
# 2. 单站测速
# ---------------------------------------------------------------------------

def test_site(host, test_file, max_bytes):
    """
    通过代理下载 test_file 的前 max_bytes 字节。
    返回 dict：status/http/ttfb_ms/bytes/elapsed_s/speed_mbps/note
    status: ok / timeout / error
    """
    # IDN 域名（如 ghf.无名氏.top）转 punycode
    try:
        host_enc = host.encode('idna').decode('ascii')
    except Exception:
        host_enc = host
    url = f'https://{host_enc}/{test_file}'

    result = {
        'host': host, 'status': 'error', 'http': '',
        'ttfb_ms': '', 'bytes': 0, 'elapsed_s': '', 'speed_mbps': 0.0, 'note': '',
    }
    start = time.perf_counter()
    try:
        with http_get(
            url,
            headers={'User-Agent': UA, 'Accept': '*/*'},
            timeout=(CONNECT_TIMEOUT, READ_TIMEOUT),
            stream=True,
            allow_redirects=True,
        ) as r:
            result['http'] = r.status_code
            if r.status_code != 200:
                result['note'] = f'HTTP {r.status_code}'
                return result
            got = 0
            first_chunk_at = None
            for chunk in r.iter_content(chunk_size=65536):
                if not chunk:
                    continue
                if first_chunk_at is None:
                    first_chunk_at = time.perf_counter()
                    result['ttfb_ms'] = round((first_chunk_at - start) * 1000)
                got += len(chunk)
                if got >= max_bytes:
                    break
                if time.perf_counter() - start > TOTAL_TIMEOUT:
                    result['status'] = 'timeout'
                    result['bytes'] = got
                    result['note'] = f'总耗时超过 {TOTAL_TIMEOUT}s，跳过'
                    return result
            elapsed = time.perf_counter() - start
            result['bytes'] = got
            result['elapsed_s'] = round(elapsed, 2)
            if got > 0 and elapsed > 0:
                result['speed_mbps'] = round(got / elapsed / 1024 / 1024, 3)
            if got > 0:
                result['status'] = 'ok'
                if got < max_bytes:
                    result['note'] = f'文件不足 {max_bytes // 1024 // 1024}MB，仅 {got / 1024 / 1024:.1f}MB'
            else:
                result['note'] = '无数据返回'
    except requests.exceptions.ConnectTimeout:
        result['status'] = 'timeout'
        result['note'] = '连接超时'
    except requests.exceptions.ReadTimeout:
        result['status'] = 'timeout'
        result['note'] = '读取超时'
    except requests.exceptions.SSLError as e:
        result['note'] = f'SSL 错误: {type(e).__name__}'
    except requests.exceptions.ConnectionError as e:
        result['note'] = f'连接失败: {type(e).__name__}'
    except Exception as e:
        result['note'] = f'{type(e).__name__}: {e}'
    return result


# ---------------------------------------------------------------------------
# 3. Excel 输出
# ---------------------------------------------------------------------------

def write_excel(results, test_file, size_mb, out_path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = '测速结果'

    headers = ['排名', '站点', '来源', '状态', 'HTTP', '首字节延迟(ms)',
               f'下载量(MB/{size_mb}MB)', '耗时(s)', '速度(MB/s)', '备注']
    ws.append(headers)

    head_fill = PatternFill('solid', fgColor='2F5496')
    head_font = Font(bold=True, color='FFFFFF')
    for col in range(1, len(headers) + 1):
        c = ws.cell(row=1, column=col)
        c.fill = head_fill
        c.font = head_font
        c.alignment = Alignment(horizontal='center', vertical='center')

    fill_ok = PatternFill('solid', fgColor='C6EFCE')
    fill_timeout = PatternFill('solid', fgColor='FFEB9C')
    fill_err = PatternFill('solid', fgColor='FFC7CE')
    status_text = {'ok': '可用', 'timeout': '超时跳过', 'error': '失败'}

    rank = 0
    for r in results:
        rank += 1 if r['status'] == 'ok' else 0
        ws.append([
            rank if r['status'] == 'ok' else '',
            r['host'],
            '、'.join(sorted(r['sources'])),
            status_text.get(r['status'], r['status']),
            r['http'],
            r['ttfb_ms'],
            round(r['bytes'] / 1024 / 1024, 2) if r['bytes'] else 0,
            r['elapsed_s'],
            r['speed_mbps'] if r['status'] == 'ok' else '',
            r['note'],
        ])
        row = ws.max_row
        fill = fill_ok if r['status'] == 'ok' else (fill_timeout if r['status'] == 'timeout' else fill_err)
        for col in range(1, len(headers) + 1):
            ws.cell(row=row, column=col).fill = fill

    widths = [6, 32, 20, 10, 8, 14, 14, 10, 12, 36]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f'A1:J{ws.max_row}'

    # 汇总 sheet
    ws2 = wb.create_sheet('测试信息')
    ok_n = sum(1 for r in results if r['status'] == 'ok')
    to_n = sum(1 for r in results if r['status'] == 'timeout')
    err_n = sum(1 for r in results if r['status'] == 'error')
    info = [
        ('测试时间', datetime.now().strftime('%Y-%m-%d %H:%M:%S')),
        ('测速文件', test_file),
        ('单站下载量上限', f'{size_mb} MB'),
        ('连接/读取/总超时', f'{CONNECT_TIMEOUT}s / {READ_TIMEOUT}s / {TOTAL_TIMEOUT}s'),
        ('站点总数', len(results)),
        ('可用', ok_n),
        ('超时跳过', to_n),
        ('失败', err_n),
        ('来源1', SOURCES['qqday']),
        ('来源2', SOURCES['akams']),
        ('来源3', SOURCES['moretools']),
    ]
    for k, v in info:
        ws2.append([k, v])
    ws2.column_dimensions['A'].width = 22
    ws2.column_dimensions['B'].width = 100

    wb.save(out_path)


# ---------------------------------------------------------------------------
# main
# ---------------------------------------------------------------------------

def main():
    ap = argparse.ArgumentParser(description='GitHub 下载代理站点测速（三来源抓取+去重+20MB 下载测速+Excel 输出）')
    ap.add_argument('--file', default=DEFAULT_TEST_FILE, help='测速用 GitHub 文件 URL')
    ap.add_argument('--size', type=int, default=20, help='单站下载量上限 MB（默认 20）')
    ap.add_argument('--workers', type=int, default=8, help='并发线程数（默认 8）')
    ap.add_argument('--limit', type=int, default=0, help='只测前 N 个站点（0=全部，调试用）')
    ap.add_argument('-o', '--output', default='', help='输出 xlsx 路径（默认按时间戳命名）')
    args = ap.parse_args()

    if not args.output:
        args.output = f'github_proxy_test_{datetime.now():%Y%m%d_%H%M%S}.xlsx'

    max_bytes = args.size * 1024 * 1024
    log(f'[配置] 测速文件: {args.file}')
    log(f'[配置] 单站下载上限: {args.size}MB, 并发: {args.workers}, 输出: {args.output}')

    site_sources = collect_sites()
    if not site_sources:
        log('未抓取到任何站点，退出')
        sys.exit(1)
    if args.limit > 0:
        site_sources = dict(list(site_sources.items())[:args.limit])
        log(f'[调试] 仅测试前 {len(site_sources)} 个站点')

    results = []
    total = len(site_sources)
    done = 0
    with ThreadPoolExecutor(max_workers=args.workers) as pool:
        futs = {
            pool.submit(test_site, host, args.file, max_bytes): host
            for host in site_sources
        }
        for fut in as_completed(futs):
            host = futs[fut]
            r = fut.result()
            r['sources'] = site_sources[host]
            results.append(r)
            done += 1
            tag = {'ok': 'OK ', 'timeout': 'TMO ', 'error': 'ERR '}[r['status']]
            speed = f"{r['speed_mbps']:.2f}MB/s" if r['status'] == 'ok' else r['note']
            log(f'[{done}/{total}] {tag} {host:<38} {speed}')

    # 排序：可用按速度降序，其后超时，最后失败
    order = {'ok': 0, 'timeout': 1, 'error': 2}
    results.sort(key=lambda r: (order[r['status']], -r['speed_mbps']))

    write_excel(results, args.file, args.size, args.output)
    ok_n = sum(1 for r in results if r['status'] == 'ok')
    log(f'[完成] 可用 {ok_n}/{total}，结果已写入 {args.output}')


if __name__ == '__main__':
    main()
